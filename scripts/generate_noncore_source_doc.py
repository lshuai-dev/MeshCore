from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "软著申报材料" / "MeshCore便携式网状通信终端固件软件V1.16.0.3_alpha_源程序鉴别材料_非核心代码.docx"
NAME = "MeshCore便携式网状通信终端固件软件 V1.16.0.3_alpha"
LINES_PER_PAGE = 50
HALF_LINES = 1500

# 仅选择界面视觉、显示驱动、输入适配和通用交互代码。
# 明确不选择 Mesh、Packet、Identity、Dispatcher、DataStore、无线业务、
# 消息模型、协议、密码学、地图坐标计算、授权校验及定位功耗策略文件。
FRONT_FILES = [
    "src/heltec/ui/theme/ui_theme_metrics.hpp",
    "src/heltec/ui/theme/ui_theme_metrics.cpp",
    "src/heltec/ui/theme/ui_theme_mono_flat.hpp",
    "src/heltec/ui/theme/ui_theme_mono_flat.cpp",
    "src/heltec/ui/theme/ui_theme_pixel.hpp",
    "src/heltec/ui/theme/ui_theme_pixel.cpp",
    "src/heltec/ui/widgets/button_roller.hpp",
    "src/heltec/ui/widgets/button_roller.cpp",
    "src/heltec/ui/widgets/top_pane.hpp",
    "src/heltec/ui/widgets/top_pane.cpp",
    "src/heltec/ui/core/focus_key_mapper.hpp",
    "src/heltec/ui/core/focus_key_mapper.cpp",
    "src/heltec/ui/core/input_types.hpp",
    "src/heltec/ui/core/input_adapter.hpp",
    "src/heltec/ui/core/input_adapter.cpp",
    "src/heltec/ui/core/input_dispatcher.hpp",
    "src/heltec/ui/core/input_dispatcher.cpp",
]

REAR_FILES = [
    "src/heltec/drivers/display/display_port.hpp",
    "src/heltec/drivers/display/display_port.cpp",
    "src/heltec/drivers/display/lv_mem_psram.h",
    "src/heltec/drivers/display/lv_mem_psram.cpp",
    "src/heltec/drivers/display/spi1_bus_lock.hpp",
    "src/heltec/drivers/display/ssd1306_port.hpp",
    "src/heltec/drivers/display/ssd1306_port.cpp",
    "src/heltec/drivers/display/st7735_port.hpp",
    "src/heltec/drivers/display/st7735_port.cpp",
    "src/heltec/drivers/display/st7789_port.hpp",
    "src/heltec/drivers/display/st7789_port.cpp",
    "src/heltec/drivers/input/key_input.hpp",
    "src/heltec/drivers/input/key_input.cpp",
    "src/heltec/drivers/input/momentary_button_detail.hpp",
    "src/heltec/drivers/input/momentary_button.hpp",
    "src/heltec/drivers/input/momentary_button.cpp",
    "src/heltec/drivers/input/touch_cal.hpp",
    "src/heltec/drivers/input/touch_gesture_input.hpp",
    "src/heltec/drivers/input/touch_gesture_input.cpp",
    "src/heltec/drivers/input/touch_input.hpp",
    "src/heltec/drivers/input/touch_input.cpp",
    "src/heltec/drivers/input/touch_port.hpp",
    "src/heltec/drivers/input/touch_port.cpp",
]

FORBIDDEN_PATH_PARTS = (
    "mesh.cpp", "packet", "identity", "datastore", "message",
    "radio", "crypto", "ed25519", "license", "map_geo", "geodesic",
    "heltecmesh", "payload", "protocol", "transport", "group_manager",
    "biz_facade", "gps", "location", "find_friend", "repeat_mode",
)


def validate_paths(files):
    for rel in files:
        low = rel.lower()
        if any(part in low for part in FORBIDDEN_PATH_PARTS):
            raise ValueError(f"禁止纳入的核心或敏感文件：{rel}")
        path = ROOT / rel
        if not path.is_file():
            raise FileNotFoundError(path)


def collect(files, limit):
    result = []
    for rel in files:
        path = ROOT / rel
        source = path.read_text(encoding="utf-8", errors="replace").splitlines()
        nonblank = [(number, line.rstrip()) for number, line in enumerate(source, 1) if line.strip()]
        if not nonblank:
            continue
        result.append((rel, 0, f"// 文件：{rel}"))
        result.append((rel, 0, f"// 【功能注释】{describe_file(rel)}"))
        emitted_from_file = 0
        for number, line in nonblank:
            # 删除源文件注释中可能存在的个人署名或具体日期；代码本身不改写。
            if re.search(r"\b(author|created by|copyright)\b", line, re.I):
                continue
            if re.search(r"\b20\d{2}[-/.年]\d{1,2}", line):
                continue
            if emitted_from_file > 0 and emitted_from_file % 80 == 0:
                result.append((rel, 0, f"// 【重要段落注释】{describe_block(rel, emitted_from_file)}"))
                if len(result) >= limit:
                    return result[:limit]
            result.append((rel, number, line))
            emitted_from_file += 1
            if len(result) >= limit:
                return result[:limit]
    if len(result) < limit:
        raise RuntimeError(f"非核心代码行数不足：需要 {limit}，实际 {len(result)}")
    return result[:limit]


def describe_file(rel):
    low = rel.lower()
    if "/theme/" in low:
        return "本文件负责界面颜色、字体、间距和控件视觉样式，不涉及通信或数据处理核心。"
    if "/widgets/" in low:
        return "本文件实现顶部状态区域或通用选择控件，用于显示和焦点交互。"
    if "focus_key_mapper" in low:
        return "本文件将物理按键动作映射为界面焦点操作，便于不同设备统一导航。"
    if "/ui/core/input" in low:
        return "本文件负责把按键或触摸输入转换为通用界面事件。"
    if "/drivers/display/" in low:
        return "本文件完成显示器初始化、刷新或显存管理等硬件适配工作。"
    if "/drivers/input/" in low:
        return "本文件完成按键、触摸或手势输入采集与界面事件上报。"
    return "本文件属于通用界面与设备适配代码，不包含软件核心通信实现。"


def describe_block(rel, offset):
    low = rel.lower()
    if "/theme/" in low:
        return f"以下代码继续处理控件样式状态和尺寸适配，当前文件内已整理约{offset}行。"
    if "/widgets/" in low:
        return f"以下代码继续处理控件创建、布局、焦点和刷新逻辑，当前文件内已整理约{offset}行。"
    if "/drivers/display/" in low:
        return f"以下代码继续处理显示接口时序、屏幕刷新或缓冲区操作，当前文件内已整理约{offset}行。"
    if "/drivers/input/" in low:
        return f"以下代码继续处理输入状态判断、消抖、坐标或手势转换，当前文件内已整理约{offset}行。"
    return f"以下代码继续处理通用界面事件和焦点分发，当前文件内已整理约{offset}行。"


def set_run(run, name, size, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_run(run, "宋体", 5)


def add_code_line(doc, sequence, original_line, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(7)
    p.paragraph_format.keep_together = True
    p.paragraph_format.widow_control = False
    prefix = f"{sequence:04d} " if original_line == 0 else f"{sequence:04d} [{original_line:04d}] "
    run = p.add_run(prefix + text.replace("\t", "  "))
    set_run(run, "宋体", 5)


def build():
    validate_paths(FRONT_FILES + REAR_FILES)
    front = collect(FRONT_FILES, HALF_LINES)
    rear = collect(REAR_FILES, HALF_LINES)
    selected = front + rear
    if len(selected) != 3000:
        raise RuntimeError("源码材料必须严格为3000行")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    cover = doc.sections[0]
    cover.top_margin = Cm(2.54)
    cover.bottom_margin = Cm(2.54)
    cover.left_margin = Cm(3.0)
    cover.right_margin = Cm(2.5)
    cover.header.paragraphs[0].text = ""
    cover.footer.paragraphs[0].text = ""

    for _ in range(3):
        doc.add_paragraph()
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run(NAME), "宋体", 20, True)
    p.paragraph_format.space_after = Pt(28)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("源程序鉴别材料"), "宋体", 19, True)
    p.paragraph_format.space_after = Pt(60)
    for text in (
        "材料范围：前30页与后30页非核心源程序",
        "源程序行数：3000行",
        "每页行数：50行",
        "版本标识：V1.16.0.3_alpha",
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run(p.add_run(text), "宋体", 12)
        p.paragraph_format.space_after = Pt(10)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    hp = section.header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run(hp.add_run(NAME), "宋体", 5)
    section.footer.paragraphs[0].clear()
    add_page_number(section.footer.paragraphs[0])

    for index, (_, original_line, code) in enumerate(selected, 1):
        add_code_line(doc, index, original_line, code)
        if index % LINES_PER_PAGE == 0 and index != len(selected):
            doc.add_page_break()

    doc.core_properties.title = f"{NAME} 源程序鉴别材料"
    doc.core_properties.subject = "非核心源程序前后各30页"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.save(OUT)
    print(f"DOCX={OUT}")
    print(f"CODE_LINES={len(selected)}")
    print(f"CODE_PAGES={len(selected) // LINES_PER_PAGE}")


if __name__ == "__main__":
    build()
