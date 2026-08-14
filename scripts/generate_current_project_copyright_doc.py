from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "软著申报材料" / "MeshCore便携式网状通信终端固件软件V1.16.0.3_alpha_软件说明书.docx"
NAME = "MeshCore便携式网状通信终端固件软件 V1.16.0.3_alpha"
INCLUDE_EMBEDDED_CODE = False


def font(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold


def para(doc, text, size=10.5, bold=False, center=False, indent=True, after=6):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if indent:
        p.paragraph_format.first_line_indent = Pt(21)
    font(p.add_run(text), size, bold)
    return p


def title(doc, text, level=1):
    p = para(doc, text, {1: 16, 2: 13, 3: 11}.get(level, 11), True, False, False, 10)
    p.paragraph_format.keep_with_next = True


def shade(cell, fill="D9EAF7"):
    prop = cell._tc.get_or_add_tcPr()
    node = OxmlElement("w:shd")
    node.set(qn("w:fill"), fill)
    prop.append(node)


def table(doc, headers, rows):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = True
    for i, value in enumerate(headers):
        t.rows[0].cells[i].text = value
        shade(t.rows[0].cells[i])
        for r in t.rows[0].cells[i].paragraphs[0].runs:
            font(r, 9, True)
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = value
            for p in cells[i].paragraphs:
                p.paragraph_format.space_after = Pt(0)
                for r in p.runs:
                    font(r, 9)
    doc.add_paragraph()


def page(doc):
    doc.add_page_break()


def code_excerpt(doc, heading, relative_path, start_line, end_line, annotations):
    """插入当前项目中的真实非核心代码，并在代码前加入中文功能注释。"""
    if not INCLUDE_EMBEDDED_CODE:
        # 每个代码节选前都有一个独立分页；纯说明书模式下同时移除该分页。
        if doc.paragraphs:
            last = doc.paragraphs[-1]
            breaks = last._p.xpath('.//w:br[@w:type="page"]')
            if breaks:
                last._element.getparent().remove(last._element)
        return
    title(doc, heading, 2)
    para(doc, f"代码来源：{relative_path}（节选第 {start_line}—{end_line} 行）", 9, False, False, False, 4)
    for note in annotations:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1)
        p.paragraph_format.line_spacing = 1
        font(p.add_run(f"// 【功能注释】{note}"), 7.5, True)
    source = (ROOT / relative_path).read_text(encoding="utf-8", errors="replace").splitlines()
    for number in range(start_line, min(end_line, len(source)) + 1):
        line = source[number - 1]
        if not line.strip():
            continue
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.line_spacing = Pt(9)
        run = p.add_run(f"{number:04d}  {line.replace(chr(9), '  ')}")
        run.font.name = "Consolas"
        run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        run.font.size = Pt(7.5)
    para(doc, "以上代码仅体现界面或设备适配过程，不包含网状路由、报文处理、加密签名、身份密钥及无线协议核心实现。", 9, False, False, True, 6)


def build():
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    sec = doc.sections[0]
    sec.top_margin, sec.bottom_margin = Cm(2.54), Cm(2.54)
    sec.left_margin, sec.right_margin = Cm(3.0), Cm(2.5)
    sec.header.paragraphs[0].text = ""
    sec.footer.paragraphs[0].text = ""
    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    doc.core_properties.title = f"{NAME} 软件说明书"
    doc.core_properties.subject = "软件著作权登记说明文档"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""

    para(doc, "", after=48)
    para(doc, NAME, 22, True, True, False, 28)
    para(doc, "软件说明书", 20, True, True, False, 70)
    para(doc, "文档类别：嵌入式固件软件使用与设计说明", 12, False, True, False, 12)
    para(doc, "软件版本：V1.16.0.3_alpha", 12, False, True, False, 12)
    para(doc, "适用范围：便携式低功耗网状通信终端", 12, False, True, False, 12)

    page(doc)
    title(doc, "目录")
    chapters = [
        "1. 项目概述", "2. 项目组成与开发环境", "3. 支持的终端与固件形态", "4. 软件总体设计",
        "5. 启动与主界面", "6. 消息通信功能", "7. 最近消息与节点信息", "8. 无线状态与参数管理",
        "9. GPS定位与位置共享", "10. GPS轨迹与动态功耗管理", "11. 指南针与查找朋友",
        "12. 地图与跟踪页面", "13. 系统设置与转发模式", "14. 输入、导航与显示适配",
        "15. 连接方式与外部应用协同", "16. 数据保存与恢复", "17. 当前版本功能变更",
        "18. 编译、烧录与验收", "19. 常见问题", "20. 软件功能总结"
    ]
    for c in chapters:
        para(doc, c, 11, False, False, False, 5)

    page(doc)
    title(doc, "1. 项目概述")
    para(doc, f"{NAME} 是当前固件仓库中的便携式网状通信终端软件。项目使用 C++ 开发，面向资源受限的微控制器和远距离低功耗无线设备，提供无需互联网和中心服务器即可运行的多跳消息通信能力。本文档依据当前工作区中的实际目录、构建配置、终端界面和版本提交记录整理。")
    para(doc, "当前版本在基础网状通信能力之上，重点实现了适合带屏便携终端的本地交互界面，并完成多种终端显示、按键、触摸、定位、方向、蜂鸣、电源和存储能力的适配。用户可以在设备端查看节点状态、发送消息、浏览最近消息、检查无线质量、管理定位与位置共享、使用地图和方向辅助，并调整系统功能。")
    title(doc, "1.1 主要使用场景", 2)
    table(doc, ["场景", "项目实际能力"], [
        ["户外联络", "终端之间通过远距离无线网络交换个人消息和群组消息。"],
        ["无公网区域", "基础通信不依赖互联网或中心服务器。"],
        ["移动位置协作", "支持定位状态、位置共享、轨迹记录、方向和目标查找。"],
        ["临时覆盖扩展", "终端可按当前版本提供的受限频率选项启用转发模式。"],
        ["外部应用接入", "固件可通过蓝牙、USB或无线局域网形态与配套客户端协同。"]
    ])

    page(doc)
    title(doc, "2. 项目组成与开发环境")
    para(doc, "当前项目采用 PlatformIO 组织多平台嵌入式构建。根目录的统一配置负责加载各终端变体，具体硬件能力、依赖和编译目标在 variants 目录中按设备划分。公共网状通信能力位于 src 目录；终端定制应用、界面、驱动、地图和配置管理集中在终端应用相关目录；可构建的不同固件角色位于 examples 目录。")
    table(doc, ["项目区域", "当前用途"], [
        ["src", "公共通信、数据包、身份、调度以及终端应用功能。"],
        ["src/终端应用", "图形界面、业务页面、提示层、导航、地图、配置和设备适配。"],
        ["variants", "不同终端型号的硬件定义和 PlatformIO 构建环境。"],
        ["examples", "伴侣终端、中继、房间服务、传感、终端聊天和调制解调器等固件入口。"],
        ["arch 与 lib", "平台适配组件和项目随附的第三方库。"],
        ["docs", "通信说明、命令说明、数据格式和使用说明等已有项目文档。"],
        ["scripts 与 build.sh", "构建辅助、补丁处理和发布文件生成。"]
    ])
    title(doc, "2.1 开发工具", 2)
    para(doc, "项目推荐使用 Visual Studio Code 与 PlatformIO 进行配置、编译和烧录。主框架为 Arduino 兼容嵌入式框架；当前配置覆盖 ESP32、nRF52、RP2040 和 STM32 等平台。无线功能使用兼容的射频驱动组件，图形终端使用轻量级嵌入式图形库完成页面和控件显示。")

    page(doc)
    title(doc, "3. 支持的终端与固件形态")
    para(doc, "当前仓库包含大量终端构建环境。与本说明书所述定制终端功能相关的设备系列包括 T096、T1、T114、T190、Tracker、Tracker V2、V2、V3、V4、V4 R8、CT62、E213、E290、Mesh Solar 和 Wireless Paper 等。不同设备的屏幕、输入、定位和存储能力不同，因此最终菜单会按编译配置显示或隐藏。")
    table(doc, ["固件形态", "项目中的实际用途"], [
        ["伴侣终端", "通过蓝牙、USB或无线局域网连接外部客户端，完成消息和配置交互。"],
        ["中继节点", "持续参与无线消息转发，用于扩大网络覆盖。"],
        ["房间服务", "提供共享消息空间和相应的节点服务能力。"],
        ["终端聊天", "直接在串行终端或本地终端中完成简化聊天操作。"],
        ["传感节点", "采集环境或位置类信息并通过网络发送。"],
        ["调制解调器", "通过标准串行帧形式为主机软件提供无线收发接口。"]
    ])
    para(doc, "本说明书重点描述当前带屏便携终端的伴侣通信与本地界面能力。无显示设备或纯中继固件不会呈现全部页面，但共享相同项目中的基础无线和设备支持能力。")

    page(doc)
    title(doc, "4. 软件总体设计")
    para(doc, "当前终端软件由无线通信、业务协调、图形界面、输入适配、设备驱动、位置工具、配置存储和外部连接等部分组成。界面不直接处理底层报文，而是通过业务接口读取节点状态、联系人、消息、无线质量和位置数据，并将用户操作转换为对应业务请求。")
    table(doc, ["功能层", "实际职责", "界面表现"], [
        ["设备层", "管理显示、按键、触摸、蜂鸣、振动、定位、电池和外设。", "设备状态和输入反馈。"],
        ["通信层", "完成无线收发、节点识别、多跳传递和消息队列处理。", "连接、消息和无线质量状态。"],
        ["业务层", "协调联系人、群组、位置共享、转发模式和设置保存。", "页面数据与操作结果。"],
        ["界面层", "组织页面、导航窗格、弹出提示、键盘和列表。", "Home等实际页面。"],
        ["存储层", "保存节点偏好、开关、周期、无线和设备设置。", "重启后恢复配置。"]
    ])
    para(doc, "界面采用页面与覆盖层相结合的结构。页面用于长期显示某类信息；覆盖层用于消息发送、文本键盘、预览、警告、校准、快捷发送、无线同步和转发模式确认。覆盖层关闭后返回原页面，减少用户在多级菜单中的跳转。")
    page(doc)
    code_excerpt(doc, "4.1 界面度量配置代码", "src/heltec/ui/theme/ui_theme_metrics.hpp", 1, 65, [
        "该部分定义不同屏幕共用的界面尺寸配置，是纯界面表现代码。",
        "度量接口把顶部区域、菜单、按键滚轮等控件尺寸集中管理，便于多设备适配。"
    ])

    page(doc)
    title(doc, "5. 启动与主界面")
    para(doc, "设备启动后首先显示启动画面，并在启动阶段准备显示、输入、无线、配置和可选传感器。启动完成后进入 Home 页面。当前版本启动画面显示固件版本，默认版本标识为 v1.16.0.3_alpha。")
    title(doc, "5.1 Home页面", 2)
    para(doc, "Home 页面显示本机节点标识、消息数量和运行状态摘要，是其他功能的起点。顶部区域用于显示当前页面标题及必要状态；导航窗格用于在 Home、Recent、Radio、Compass、Find Friend、GPS、Tracker 和 System 等页面之间切换。")
    table(doc, ["显示项", "含义"], [
        ["节点标识", "用于识别当前终端的公开标识摘要。"],
        ["消息数量", "表示当前可见或待处理消息的数量状态。"],
        ["状态文字", "显示网络、设备或操作相关的简短状态。"],
        ["顶部标题", "随当前页面或导航预览同步变化。"],
        ["导航图标", "对应当前终端实际启用的功能页面。"]
    ])
    title(doc, "5.2 启动检查", 2)
    para(doc, "使用前应确认设备已连接匹配的天线、电量足够并能正常进入 Home 页面。带定位功能的终端首次获取位置可能需要在室外等待。带触摸屏的设备可直接点击列表和按钮；单键或多键设备按屏幕操作提示移动焦点并确认。")
    page(doc)
    code_excerpt(doc, "5.3 Home页面创建代码", "src/heltec/ui/screens/home_screen.cpp", 43, 82, [
        "该代码创建Home页面中的节点标识、消息数量和状态标签。",
        "这里只负责读取界面快照并刷新文字，不包含消息传输和节点路由算法。"
    ])

    page(doc)
    title(doc, "6. 消息通信功能")
    para(doc, "当前项目支持个人联系人消息和群组消息。消息发送覆盖层可先选择消息类别，再选择具体联系人或群组，随后从预设消息或自定义文本中选择内容。支持屏幕键盘的设备可以直接输入自定义文本。")
    table(doc, ["步骤", "当前软件操作", "结果"], [
        ["1", "打开消息发送功能。", "显示个人与群组等目标类别。"],
        ["2", "选择联系人或群组。", "界面保存本次发送目标。"],
        ["3", "选择预设文本或打开自定义输入。", "显示待发送内容。"],
        ["4", "确认发送。", "终端提交消息并显示成功或失败提示。"],
        ["5", "返回 Recent 或消息预览。", "查看近期收到的消息摘要。"]
    ])
    title(doc, "6.1 Quick Ping", 2)
    para(doc, "项目提供 Quick Ping 快捷操作层，用于减少常用短消息的操作步骤。当前版本优化了该区域的背景和焦点颜色，使按键设备更容易识别当前选项。快捷消息仍需要选定目标，发送结果由提示层反馈。")
    title(doc, "6.2 消息接收提示", 2)
    para(doc, "收到个人或群组文本后，软件将消息加入当前内存预览队列并触发界面提示。预览区域可保留近期消息，支持单击或方向键切换，并显示消息接收后经过的时间。该机制用于快速查看，不等同于无限期的文件化聊天记录。")
    page(doc)
    code_excerpt(doc, "6.3 消息选择界面代码", "src/heltec/ui/overlays/touch_send_message_overlay.cpp", 85, 134, [
        "该部分只创建消息发送页面的标题、列表、底部提示和触摸行。",
        "真正的消息编码、路由和无线发送过程未纳入本文档。"
    ])

    page(doc)
    title(doc, "7. 最近消息与节点信息")
    para(doc, "Recent 页面在当前版本中主要显示最近接收到的节点广播信息。页面最多维护 16 条近期节点信息，并依据不同屏幕高度动态计算一页可显示的行数。用户可通过按键翻页、自动滚动或触摸拖动浏览。")
    table(doc, ["区域", "当前行为"], [
        ["节点列表", "展示近期接收到的节点名称或可识别摘要。"],
        ["分页", "依据屏幕高度复用现有行控件进行翻页。"],
        ["焦点", "所有可操作列表行加入统一焦点组。"],
        ["触摸滚动", "支持图形库原生拖动滚动。"],
        ["消息预览", "由独立的近期消息覆盖层负责，不再依赖旧的文件化历史模块。"]
    ])
    para(doc, "用户查看 Recent 页面时，可使用上下或左右输入切换项目。列表项目超出显示区域后，当前焦点会自动滚动到可见范围。触摸设备可直接拖动列表并点击目标。当前页面强调近期网络节点发现，不作为完整的长期通信记录存档。")
    page(doc)
    code_excerpt(doc, "7.1 Recent页面控件代码", "src/heltec/ui/screens/recent_screen.cpp", 1, 62, [
        "该节选展示Recent页面使用的文本格式、行状态和页面初始化辅助逻辑。",
        "代码用于近期信息的界面呈现，不涉及无线广播解析的核心过程。"
    ])

    page(doc)
    title(doc, "8. 无线状态与参数管理")
    para(doc, "Radio 页面用于查看当前无线工作状态和最近一次接收质量。当前版本新增最近一次接收消息的 RSSI 与 SNR 显示，便于用户判断现场信号强弱和接收质量。页面同时提供当前无线配置的可读摘要。")
    table(doc, ["显示内容", "项目中的含义"], [
        ["无线状态", "反映无线模块是否处于可用状态。"],
        ["RSSI", "最近一次接收信号的强度参考。"],
        ["SNR", "最近一次接收信号相对于噪声的质量参考。"],
        ["工作方案", "显示当前使用的无线区域或预设配置。"],
        ["同步提示", "无线参数发生变更时显示应用或同步状态。"]
    ])
    para(doc, "项目支持通过预设方式选择无线工作方案，并在设置变更后更新无线模块。用户应选择符合终端硬件和所在地无线电管理要求的配置。为保护项目核心实现，本说明书不列出内部报文、计算过程和完整底层参数表。")
    page(doc)
    code_excerpt(doc, "8.1 Radio状态页面代码", "src/heltec/ui/screens/radio_screen.cpp", 1, 72, [
        "该代码负责创建Radio页面标签并刷新RSSI、SNR等用户可见状态。",
        "节选不包含无线驱动初始化、调制参数应用或报文收发实现。"
    ])

    page(doc)
    title(doc, "9. GPS定位与位置共享")
    para(doc, "GPS 页面是当前版本新增并集中整理的定位功能页面。此前分散在 System 页面中的 GPS Track 和 Location Share 设置已迁移到该页面。页面可显示 GPS 开关、位置共享开关、广播周期、轨迹开关、定位状态、卫星信息、经纬度、高度和实时移动速度。")
    table(doc, ["页面项目", "实际功能"], [
        ["GPS", "启用或停用定位模块。"],
        ["Location Share", "控制是否按所选周期共享本机位置。"],
        ["Adv Interval", "选择位置共享的广播间隔。"],
        ["GPS Track", "控制本机轨迹记录功能。"],
        ["Fix与卫星", "显示当前是否定位及卫星相关状态。"],
        ["经纬度、高度、速度", "显示当前有效定位数据和实时移动速度。"]
    ])
    para(doc, "开启 GPS Track 或 Location Share 时，软件会自动开启 GPS，避免功能开关与定位模块状态不一致。经纬度手工输入完成后，当前版本优化了页面返回和数据同步流程。位置数据属于敏感信息，启用共享前应确认使用场景和接收范围。")
    page(doc)
    code_excerpt(doc, "9.1 GPS页面开关与标签代码", "src/heltec/ui/screens/gps_screen.cpp", 60, 128, [
        "该节选创建GPS开关、位置共享开关和广播周期选择控件。",
        "代码属于页面布局与事件绑定部分，未展示定位算法、位置广播和功耗决策核心。"
    ])

    page(doc)
    title(doc, "10. GPS轨迹与动态功耗管理")
    para(doc, "当前版本加入 GPS 动态功耗管理。软件根据外部供电状态、屏幕状态、当前页面、轨迹任务、位置共享周期和临时定位发送窗口决定 GPS 的工作状态，在定位及时性和电池续航之间进行协调。")
    table(doc, ["使用状态", "当前版本的GPS策略"], [
        ["GPS页面或地图页面正在使用", "保持较积极的定位状态，便于及时刷新。"],
        ["GPS轨迹开启", "为轨迹任务维持所需定位能力。"],
        ["短周期位置共享", "保持能够满足频繁广播的位置更新。"],
        ["长周期位置共享", "在接近广播前的时间窗口提高定位状态，其余时间降低功耗。"],
        ["熄屏持续约30秒且无高优先级任务", "允许定位模块进入较低功耗状态。"],
        ["Find Friend开启且熄屏", "允许在满足功能需求的前提下降低GPS功耗。"]
    ])
    para(doc, "动态功耗管理仅在设备具备相应定位和电源管理能力时生效。外部供电设备与电池供电设备的策略可能不同。用户打开地图、GPS页面或执行位置发送时，软件会优先保证当前任务所需的定位响应。")

    page(doc)
    title(doc, "11. 指南针与查找朋友")
    para(doc, "Compass 页面用于显示终端方向信息。当前版本调整了页面布局，避免航向数值被其他内容遮挡。设备具备方向传感器时，用户可在校准后查看航向；不具备相应硬件的终端不会提供完整指南针能力。")
    title(doc, "11.1 Find Friend页面", 2)
    para(doc, "Find Friend 已由系统设置中的分散选项调整为独立页面，并新增明确的功能启用开关。启用状态可保存，设备重启后按项目设置恢复。页面用于选择目标并结合本机与目标的可用位置显示方向或距离辅助。")
    table(doc, ["操作", "说明"], [
        ["启用功能", "打开 Find Friend 开关并保存状态。"],
        ["选择目标", "从当前可用联系人中选择需要查找的对象。"],
        ["查看信息", "读取目标名称、方向、距离或位置状态。"],
        ["更新判断", "等待新的位置或无线信息后刷新。"],
        ["退出", "返回导航页面，不影响已保存的启用状态。"]
    ])
    para(doc, "T1 类小屏终端的 Find Friend 页面已进行专门布局优化，主要信息可在一屏内显示。该功能依赖有效位置数据和消息更新，不应作为专业搜救或生命安全场景中的唯一依据。")

    page(doc)
    title(doc, "12. 地图与跟踪页面")
    para(doc, "带地图能力的构建会显示 Tracker 页面。项目中的地图功能由地图面板、本地地图资源、地图偏好和地理坐标换算等部分组成，可显示当前位置、目标或轨迹类信息。地图内容来自设备可访问的本地资源，而不是必须依赖在线地图。")
    table(doc, ["地图功能", "用户可见结果"], [
        ["地图加载", "显示当前本地地图资源覆盖的区域。"],
        ["当前位置", "定位有效时在地图中显示本机位置。"],
        ["移动与缩放", "通过按键或触摸浏览相邻区域及不同范围。"],
        ["目标辅助", "结合联系人位置或跟踪状态显示目标信息。"],
        ["轨迹相关", "在启用GPS Track时为移动过程提供可视化辅助。"]
    ])
    para(doc, "地图页面在打开期间会影响 GPS 动态功耗策略，以保证位置刷新。地图资源不存在、当前位置超出资源范围或存储不可用时，页面可能无法显示完整地图。项目中的内部地图格式、索引和坐标处理细节不在本登记说明文档中披露。")

    page(doc)
    title(doc, "13. 系统设置与转发模式")
    para(doc, "System 页面集中管理终端通用设置。当前页面采用可滚动设置列表，并根据设备能力显示开关、下拉选项和操作行。当前版本恢复使用原生下拉列表，按键可以在展开状态中循环选择，触摸设备可直接点击选项。")
    table(doc, ["设置类型", "项目中的实际内容"], [
        ["连接设置", "管理伴侣连接等可用开关。"],
        ["显示和交互", "管理屏幕、背光、导航和提示相关行为。"],
        ["声音", "按设备能力调整蜂鸣器启用状态或音量等级。"],
        ["无线辅助", "启用转发模式并选择允许的转发频率。"],
        ["数据操作", "提供清除数据和恢复默认设置入口。"],
        ["设备信息", "查看版本和当前终端相关信息。"]
    ])
    title(doc, "13.1 转发模式", 2)
    para(doc, "当前版本新增 Enable Repeat Mode 开关。用户打开后，从项目限定的 433、869 或 918 MHz 选项中选择转发频率；选择期间按 ESC 取消时保持转发模式关闭；启用成功后界面给出状态提示。该功能只允许当前软件认可的选项，并应在符合所在地规定的前提下使用。")
    page(doc)
    code_excerpt(doc, "13.2 系统设置行创建代码", "src/heltec/ui/screens/system_screen_rows.cpp", 90, 174, [
        "该代码用于创建系统设置中的操作行、开关行和下拉选择行。",
        "节选重点说明控件布局和事件绑定，不包含转发模式的无线应用逻辑。"
    ])

    page(doc)
    title(doc, "14. 输入、导航与显示适配")
    para(doc, "项目针对不同终端采用统一输入事件和焦点管理。实际输入来源包括单键、多键、旋钮、触摸点击、触摸拖动和手势。输入适配层将不同硬件动作转换为上一项、下一项、确认、返回和快捷操作，使同一套页面能够运行在多种屏幕和输入组合上。")
    table(doc, ["输入方式", "当前项目行为"], [
        ["方向键或旋钮", "移动焦点、切换页面或改变列表选项。"],
        ["确认键", "进入页面、打开控件或提交操作。"],
        ["ESC/返回", "关闭弹出层、取消选择或返回上一状态。"],
        ["触摸点击", "选中列表行、开关或操作按钮。"],
        ["触摸拖动", "滚动长列表。"],
        ["触摸手势", "在支持设备上执行页面或快捷交互。"]
    ])
    para(doc, "当前版本统一将可操作列表行加入焦点组，解决小屏页面项目超出可视区域后的键盘翻页和自动滚动问题。V4 R8 TFT 对列表行高度、间距和选中状态进行了单独优化；导航图标、选中状态及布局也会按设备屏幕类型变化。")
    page(doc)
    code_excerpt(doc, "14.1 按键焦点映射代码", "src/heltec/ui/core/focus_key_mapper.cpp", 1, 72, [
        "该代码把设备按键转换为上一项、下一项、确认和返回等通用界面动作。",
        "映射层只处理人机交互，不读取或修改无线通信数据。"
    ])
    page(doc)
    code_excerpt(doc, "14.2 导航页面配置代码", "src/heltec/ui/app/ui_app_navigation.cpp", 96, 138, [
        "该代码依据终端能力配置Home、Recent、Radio、Compass、GPS、Tracker和System页面。",
        "条件配置确保没有相应硬件能力的设备不会显示不可用页面。"
    ])
    page(doc)
    code_excerpt(doc, "14.3 显示端口初始化代码", "src/heltec/drivers/display/display_port.cpp", 1, 64, [
        "该代码根据编译配置选择匹配的屏幕端口并执行界面显示初始化。",
        "显示驱动仅负责屏幕输出，不包含项目的通信和安全核心。"
    ])

    page(doc)
    title(doc, "15. 连接方式与外部应用协同")
    para(doc, "伴侣终端固件可按构建环境启用蓝牙、USB或无线局域网连接。外部客户端通过所选连接方式读取节点状态、联系人、群组、消息和配置，并向终端提交发送或设置请求。当前定制终端默认允许伴侣连接，具体可用连接方式由编译目标和硬件能力决定。")
    table(doc, ["连接方式", "适用特点"], [
        ["蓝牙", "适合移动设备近距离无线连接。"],
        ["USB", "适合有线配置、串行通信和稳定供电场景。"],
        ["无线局域网", "适合具备相应网络能力的终端构建。"],
        ["本地屏幕", "无需外部客户端即可查看状态和执行常用操作。"]
    ])
    para(doc, "当前版本修复了重新刷写固件后蓝牙默认状态异常的问题。连接失败时，应先确认使用了对应连接形态的固件、终端连接功能已开启、外部设备权限正常，并避免多个客户端同时占用同一连接。")

    page(doc)
    title(doc, "16. 数据保存与恢复")
    para(doc, "项目使用本地配置存储保存节点偏好和终端设置，包括部分无线选项、伴侣连接状态、GPS开关、位置共享周期、GPS Track、Find Friend、转发模式及其他设备设置。软件启动时会读取并校验设置，异常值会被限制或恢复为安全状态。")
    table(doc, ["数据类别", "保存行为"], [
        ["节点偏好", "用于恢复终端身份显示和常用工作选项。"],
        ["GPS与位置共享", "保存开关及位置共享周期。"],
        ["Find Friend", "保存功能是否启用。"],
        ["转发模式", "仅在频率选项有效时恢复，否则保持关闭。"],
        ["伴侣连接", "保存允许的连接状态。"],
        ["界面与设备选项", "按具体终端能力保存。"]
    ])
    para(doc, "System 页面提供 Clear Data 和 Factory Reset 等数据操作入口。执行前应确认不再需要联系人、设置或本地信息，并在操作过程中保持供电稳定。当前版本已移除旧的文件化 MessageHistory 功能，近期消息主要通过运行期间的预览队列展示。")

    page(doc)
    title(doc, "17. 当前版本功能变更")
    para(doc, "根据当前仓库 v1.16.0.2 以来的项目提交记录，v1.16.0.3_alpha 主要围绕定位、低功耗、目标查找、转发模式、近期信息、无线质量和多设备交互适配进行了集中修改。")
    table(doc, ["变更分类", "当前版本实际修改"], [
        ["GPS", "新增独立GPS页面、实时速度、轨迹与位置共享集中管理。"],
        ["功耗", "增加GPS动态功耗管理和熄屏后的低功耗切换。"],
        ["Find Friend", "迁移为独立页面，增加启用开关、持久化和小屏布局优化。"],
        ["转发模式", "增加系统开关、限定频率选择、取消处理和成功提示。"],
        ["Recent", "恢复近期节点广播显示，最多16条，按屏幕动态分页。"],
        ["消息预览", "使用独立32条近期消息队列，支持切换和动态时长显示。"],
        ["Radio", "新增最近一次接收的RSSI和SNR。"],
        ["交互", "统一焦点滚动，优化触摸列表、导航、Quick Ping和不同屏幕布局。"],
        ["设备适配", "完善Tracker V2资源和构建配置，优化V4 R8 TFT显示。"],
        ["问题修复", "修复刷写后蓝牙默认状态及经纬度输入后的返回同步问题。"]
    ])

    page(doc)
    title(doc, "18. 编译、烧录与验收")
    para(doc, "开发人员在 PlatformIO 中选择与终端型号及固件形态一致的构建环境进行编译。项目根配置会加载 variants 下的设备配置，并按环境组合公共源文件、终端应用、显示驱动、输入驱动和可选传感器。生成的固件只能烧录到匹配的设备型号。")
    table(doc, ["验收编号", "检查内容", "预期结果"], [
        ["T01", "启动与版本", "正常进入Home页面并显示v1.16.0.3_alpha。"],
        ["T02", "页面导航", "可进入当前设备启用的各个页面并正常返回。"],
        ["T03", "消息发送", "可选择个人或群组目标并获得发送结果提示。"],
        ["T04", "近期信息", "Recent显示近期节点，预览层显示近期消息。"],
        ["T05", "无线状态", "Radio可显示当前状态及最近RSSI、SNR。"],
        ["T06", "GPS", "支持设备可显示定位、位置共享、轨迹和速度。"],
        ["T07", "Find Friend", "开关可保存并能选择目标查看辅助信息。"],
        ["T08", "转发模式", "仅允许规定选项，取消和启用提示正常。"],
        ["T09", "触摸或按键", "列表焦点、滚动、确认和返回行为正常。"],
        ["T10", "重启恢复", "已保存设置按设计恢复，异常配置不被错误启用。"]
    ])

    page(doc)
    title(doc, "19. 常见问题")
    table(doc, ["现象", "结合当前项目的检查方法"], [
        ["无法进入某个页面", "确认所选终端是否编译启用了GPS、Compass、Map或触摸等能力。"],
        ["消息无法发送", "检查天线、无线状态、目标、信号质量和两端配置。"],
        ["Recent没有消息内容", "Recent主要显示近期节点；消息内容应在消息预览层查看。"],
        ["GPS没有定位", "打开GPS并在室外等待，检查设备是否具有定位模块。"],
        ["Find Friend无方向", "检查本机与目标是否都有可用位置，必要时完成指南针校准。"],
        ["转发模式无法开启", "必须选择软件允许的频率选项；取消选择时保持关闭。"],
        ["列表无法看到后续项目", "使用方向输入移动焦点，或在触摸设备上拖动列表。"],
        ["蓝牙连接异常", "确认刷入蓝牙伴侣固件、连接开关和移动设备权限。"],
        ["地图空白", "确认该构建启用地图、存储可用且已有匹配的本地地图资源。"],
        ["设置未保存", "稳定供电后重新确认设置，必要时检查存储或恢复默认设置。"]
    ])
    para(doc, "设备出现持续异常时，应记录终端型号、构建环境、固件版本、复现步骤和屏幕现象。登记说明文档不包含内部调试命令和核心代码，研发排查应在受控环境中使用项目自己的调试手段。")

    page(doc)
    title(doc, "20. 软件功能总结")
    para(doc, f"{NAME} 已在当前项目中形成面向多种便携终端的完整固件组合。软件以低功耗远距离网状通信为基础，提供本地屏幕状态查看、个人与群组消息、近期节点和消息预览、无线质量、GPS定位、位置共享、轨迹、指南针、目标查找、地图、转发模式、系统设置以及蓝牙、USB和无线局域网协同能力。")
    para(doc, "当前 v1.16.0.3_alpha 版本的主要特点是将定位功能集中到独立页面，增加GPS动态功耗管理和实时速度显示，完善Find Friend页面，加入受限转发模式，恢复近期节点显示，增加RSSI与SNR，并统一不同终端上的焦点、滚动、触摸和导航体验。")
    para(doc, "本文档内容均来源于当前工作区的项目结构、构建配置、页面组成和版本修改记录。为保护软件核心实现，文档未包含任何源程序正文、函数实现、核心算法、协议字段、密钥内容或能够还原关键技术细节的材料。")
    title(doc, "20.1 文档与软件对应关系", 2)
    table(doc, ["项目", "对应内容"], [
        ["软件名称", "MeshCore便携式网状通信终端固件软件"],
        ["软件版本", "V1.16.0.3_alpha"],
        ["项目类型", "嵌入式网状通信终端固件"],
        ["主要平台", "ESP32、nRF52、RP2040、STM32等项目已配置平台"],
        ["主要交互", "屏幕、按键、旋钮、触摸及外部伴侣客户端"],
        ["文档范围", "当前项目实际功能、操作、构建与验收说明"],
        ["排除内容", "源代码、核心实现、协议内部细节和敏感数据"]
    ])

    # 附录按非核心模块逐页展示真实代码，使说明书总页数超过60页。
    appendix_samples = [
        ("A.1 单色主题初始化", "src/heltec/ui/theme/ui_theme_mono_flat.cpp", 1, 48,
         ["该段负责单色屏幕的基础主题初始化。", "主题代码只改变控件外观，不参与通信业务。"]),
        ("A.2 单色控件样式", "src/heltec/ui/theme/ui_theme_mono_flat.cpp", 49, 102,
         ["该段设置标签、按钮和选择状态的单色显示效果。"]),
        ("A.3 单色页面样式", "src/heltec/ui/theme/ui_theme_mono_flat.cpp", 103, 160,
         ["该段根据控件类型应用页面样式和边框规则。"]),
        ("A.4 像素主题初始化", "src/heltec/ui/theme/ui_theme_pixel.cpp", 1, 52,
         ["该段创建适用于小尺寸屏幕的像素化主题。"]),
        ("A.5 像素主题控件处理", "src/heltec/ui/theme/ui_theme_pixel.cpp", 53, 112,
         ["该段处理按钮、标签和选中状态的视觉表现。"]),
        ("A.6 控件主题公共入口", "src/heltec/ui/theme/ui_widget_theme.cpp", 1, 58,
         ["该段是通用控件主题的入口和基础颜色配置。"]),
        ("A.7 控件边框与背景", "src/heltec/ui/theme/ui_widget_theme.cpp", 59, 118,
         ["该段集中设置控件背景、边框和透明度。"]),
        ("A.8 控件文字样式", "src/heltec/ui/theme/ui_widget_theme.cpp", 119, 178,
         ["该段处理文本颜色、对齐和不同状态显示。"]),
        ("A.9 按钮滚轮初始化", "src/heltec/ui/widgets/button_roller.cpp", 1, 62,
         ["该段初始化按钮滚轮使用的普通和选中样式。"]),
        ("A.10 按钮滚轮焦点", "src/heltec/ui/widgets/button_roller.cpp", 63, 126,
         ["该段根据焦点状态更新滚轮项目的显示效果。"]),
        ("A.11 按钮滚轮布局", "src/heltec/ui/widgets/button_roller.cpp", 127, 194,
         ["该段计算可见项目和滚动位置，属于界面布局逻辑。"]),
        ("A.12 顶部状态区域创建", "src/heltec/ui/widgets/top_pane.cpp", 1, 58,
         ["该段创建页面顶部状态区域使用的基础对象。"]),
        ("A.13 顶部状态文字刷新", "src/heltec/ui/widgets/top_pane.cpp", 59, 118,
         ["该段更新标题和状态文字，不处理底层业务数据。"]),
        ("A.14 顶部状态图标布局", "src/heltec/ui/widgets/top_pane.cpp", 119, 182,
         ["该段负责状态图标的排列和可见性。"]),
        ("A.15 输入类型定义", "src/heltec/ui/core/input_types.hpp", 1, 77,
         ["该段定义按键、触摸和通用界面输入的数据类型。"]),
        ("A.16 输入适配入口", "src/heltec/ui/core/input_adapter.cpp", 1, 58,
         ["该段把设备输入转换为界面能够识别的通用动作。"]),
        ("A.17 输入事件分发", "src/heltec/ui/core/input_dispatcher.cpp", 1, 43,
         ["该段将已经标准化的输入事件交给当前界面处理。"]),
        ("A.18 输入处理管线", "src/heltec/ui/core/input_pipeline.cpp", 1, 66,
         ["该段组织输入读取、转换和分发的先后顺序。"]),
        ("A.19 界面动画调度", "src/heltec/ui/core/ui_motion_scheduler.cpp", 1, 62,
         ["该段负责界面动画任务的创建和调度。"]),
        ("A.20 界面动画执行", "src/heltec/ui/core/ui_motion_scheduler.cpp", 63, 128,
         ["该段处理动画状态更新和完成回调。"]),
        ("A.21 屏幕对象公共行为", "src/heltec/ui/core/abstract_screen.cpp", 1, 62,
         ["该段提供所有功能页面共用的创建和进入退出行为。"]),
        ("A.22 屏幕焦点管理", "src/heltec/ui/core/abstract_screen.cpp", 63, 132,
         ["该段管理页面内控件的焦点注册和切换。"]),
        ("A.23 覆盖层公共行为", "src/heltec/ui/core/abstract_overlay.cpp", 1, 51,
         ["该段实现提示框、键盘和快捷窗口共用的覆盖层行为。"]),
        ("A.24 延迟界面任务", "src/heltec/ui/core/ui_deferred_queue.cpp", 1, 70,
         ["该段用于延后执行界面任务，避免在输入回调中直接修改复杂控件。"]),
        ("A.25 显示内存管理", "src/heltec/drivers/display/lv_mem_psram.cpp", 1, 40,
         ["该段为图形界面提供外部内存申请和释放适配。"]),
        ("A.26 按键输入初始化", "src/heltec/drivers/input/key_input.cpp", 1, 62,
         ["该段初始化物理按键并准备界面输入设备。"]),
        ("A.27 按键状态读取", "src/heltec/drivers/input/key_input.cpp", 63, 132,
         ["该段读取按键状态并转换为界面按键值。"]),
        ("A.28 瞬时按键消抖", "src/heltec/drivers/input/momentary_button.cpp", 1, 68,
         ["该段处理机械按键消抖和按下释放状态。"]),
        ("A.29 触摸手势识别", "src/heltec/drivers/input/touch_gesture_input.cpp", 1, 72,
         ["该段识别点击、滑动和长按等基础手势。"]),
        ("A.30 触摸坐标输入", "src/heltec/drivers/input/touch_input.cpp", 1, 82,
         ["该段读取触摸坐标并生成图形界面的指针输入。"]),
    ]
    if INCLUDE_EMBEDDED_CODE:
        page(doc)
        title(doc, "附录A 非核心源代码与注释")
        para(doc, "本附录继续展示当前项目中与界面显示、通用控件、输入适配和设备交互有关的真实源代码。每个节选均增加功能注释，用于说明代码用途。附录不包含网状路由、报文格式、消息传输、加密签名、身份密钥、授权校验及其他核心实现。")
        for index, (heading, path, start, end, notes) in enumerate(appendix_samples):
            if index > 0:
                page(doc)
            code_excerpt(doc, heading, path, start, end, notes)

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
