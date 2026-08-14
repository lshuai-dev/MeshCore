from pathlib import Path
import re
import sys

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


SOFTWARE_NAME = "便携式离网网状通信终端软件 V1.0"
ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "scripts" / "generate_software_copyright_manual.ps1"
OUTPUT_DIR = ROOT / "docs" / "软著申报材料"
DOCX = OUTPUT_DIR / "便携式离网网状通信终端软件V1.0_使用与设计说明书.docx"


def strings(line: str):
    return [s.replace("$softwareName", SOFTWARE_NAME) for s in re.findall(r'"(.*?)"', line)]


def shade(cell, fill="D9EAF7"):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False):
    cell.text = text
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        for run in paragraph.runs:
            run.bold = bold
            run.font.name = "Times New Roman"
            run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
            run.font.size = Pt(9)


def add_paragraph(doc, text, size=10, bold=False, align=0, first_indent=True, after=6):
    p = doc.add_paragraph()
    p.alignment = {0: WD_ALIGN_PARAGRAPH.LEFT, 1: WD_ALIGN_PARAGRAPH.CENTER}.get(
        align, WD_ALIGN_PARAGRAPH.LEFT
    )
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.25
    if first_indent:
        p.paragraph_format.first_line_indent = Pt(21)
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    return p


def add_title(doc, text, level=1):
    size = {1: 16, 2: 13, 3: 11}.get(level, 11)
    p = add_paragraph(doc, text, size=size, bold=True, first_indent=False, after=10)
    p.paragraph_format.keep_with_next = True


def parse_actions():
    lines = SOURCE.read_text(encoding="utf-8-sig").splitlines()
    start = next(i for i, line in enumerate(lines) if "$section.Headers.Item(1).Range.Text" in line) + 2
    end = next(i for i, line in enumerate(lines) if "$doc.Fields.Update()" in line)
    actions = []
    i = start
    while i < end:
        line = lines[i].strip()
        if line == "Add-PageBreak":
            actions.append(("break",))
        elif line.startswith("Add-Title "):
            values = strings(line)
            if values:
                level_match = re.search(r"\s(\d+)\s*$", line)
                actions.append(("title", values[0], int(level_match.group(1)) if level_match else 1))
        elif line.startswith("Add-Paragraph"):
            values = strings(line)
            if "-Text $softwareName" in line:
                values = [SOFTWARE_NAME]
            if values:
                size_match = re.search(r"-Size\s+(\d+)", line)
                align_match = re.search(r"-Alignment\s+(\d+)", line)
                after_match = re.search(r"-SpaceAfter\s+(\d+)", line)
                actions.append((
                    "paragraph",
                    values[0],
                    int(size_match.group(1)) if size_match else 10,
                    "-Bold $true" in line,
                    int(align_match.group(1)) if align_match else 0,
                    "-FirstLineIndent 0" not in line,
                    int(after_match.group(1)) if after_match else 6,
                ))
        elif line.startswith("Add-Table "):
            headers = strings(line)
            rows = []
            i += 1
            while i < end and lines[i].strip() != ")":
                row_values = strings(lines[i])
                if row_values:
                    rows.append(row_values)
                i += 1
            actions.append(("table", headers, rows))
        elif line.startswith("$tocItems = @("):
            toc = []
            i += 1
            while i < end and lines[i].strip() != ")":
                toc.extend(strings(lines[i]))
                i += 1
            actions.append(("toc", toc))
            # Skip the following PowerShell foreach; the action above renders it.
            if i + 1 < end and "foreach ($item in $tocItems)" in lines[i + 1]:
                i += 1
        i += 1
    return actions


def build():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.5)
    section.header_distance = Cm(0)
    section.footer_distance = Cm(0)

    normal = doc.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10)

    doc.core_properties.title = f"{SOFTWARE_NAME} 使用与设计说明书"
    doc.core_properties.subject = "计算机软件著作权登记材料"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.keywords = ""
    doc.core_properties.comments = ""

    for action in parse_actions():
        kind = action[0]
        if kind == "break":
            doc.add_page_break()
        elif kind == "title":
            add_title(doc, action[1], action[2])
        elif kind == "paragraph":
            _, text, size, bold, align, indent, after = action
            add_paragraph(doc, text, size, bold, align, indent, after)
        elif kind == "toc":
            for item in action[1]:
                add_paragraph(doc, item, size=11, first_indent=False, after=5)
        elif kind == "table":
            headers, rows = action[1], action[2]
            table = doc.add_table(rows=1, cols=len(headers))
            table.style = "Table Grid"
            table.autofit = True
            for idx, header in enumerate(headers):
                set_cell_text(table.rows[0].cells[idx], header, bold=True)
                shade(table.rows[0].cells[idx])
            for row in rows:
                cells = table.add_row().cells
                for idx, value in enumerate(row[: len(headers)]):
                    set_cell_text(cells[idx], value)
            doc.add_paragraph().paragraph_format.space_after = Pt(0)

    # Prevent accidental headers/footers and make page starts predictable.
    for section in doc.sections:
        section.header.paragraphs[0].text = ""
        section.footer.paragraphs[0].text = ""

    doc.save(DOCX)
    print(f"DOCX={DOCX}")


if __name__ == "__main__":
    try:
        build()
    except Exception as exc:
        print(f"ERROR={exc}", file=sys.stderr)
        raise
