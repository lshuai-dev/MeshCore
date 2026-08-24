from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "软著申报材料" / "MeshCore便携式网状通信终端固件软件V1.16.0.3_alpha_源程序鉴别材料.docx"
NAME = "MeshCore便携式网状通信终端固件软件 V1.16.0.3_alpha"
LINES_PER_PAGE = 50
TOTAL_LINES = 3000
CODE_FONT_PT = 10.5  # 中文字号“五号”对应 10.5 磅
CJK_PATTERN = re.compile(r"([\u3400-\u4dbf\u4e00-\u9fff\u3000-\u303f\uff00-\uffef]+)")

# 章节顺序与《软件说明书》完全一致。quota 包含章节、文件和功能注释行。
CHAPTERS = [
    (1, "项目概述", 80, ["src/heltec/ui/app/ui_app.hpp"]),
    (2, "项目组成与开发环境", 120, ["src/heltec/ui/app/ui_app_ids.hpp", "src/heltec/ui/core/screen_id.hpp", "src/heltec/ui/core/ht_event.h", "src/heltec/ui/core/ht_event.cpp"]),
    (3, "支持的终端与固件形态", 120, ["src/heltec/drivers/display/display_port.cpp", "src/heltec/drivers/display/display_port.hpp", "src/heltec/ui/core/board_features.cpp", "variants/heltec_e213/HeltecE213Board.cpp"]),
    (4, "软件总体设计", 180, ["src/heltec/ui/core/ui_surface.cpp", "src/heltec/ui/core/ui_surface.hpp", "src/heltec/ui/core/surface_manager.hpp", "src/heltec/ui/core/surface_manager.cpp"]),
    (5, "启动与主界面", 170, ["src/heltec/ui/overlays/splash_overlay.cpp", "src/heltec/ui/screens/home_screen.cpp", "src/heltec/ui/screens/home_screen.hpp"]),
    (6, "消息通信功能", 220, ["src/heltec/ui/overlays/send_message_model.hpp", "src/heltec/ui/overlays/send_message_model.cpp", "src/heltec/ui/overlays/send_message_overlay.cpp"]),
    (7, "最近消息与节点信息", 180, ["src/heltec/ui/screens/recent_screen.hpp", "src/heltec/ui/screens/recent_screen.cpp"]),
    (8, "无线状态与参数管理", 150, ["src/heltec/ui/screens/radio_screen.hpp", "src/heltec/ui/screens/radio_screen.cpp", "src/heltec/app/mesh_app_ui_radio.cpp"]),
    (9, "GPS定位与位置共享", 200, ["src/heltec/ui/screens/gps_screen.hpp", "src/heltec/ui/screens/gps_screen.cpp"]),
    (10, "GPS轨迹与动态功耗管理", 120, ["src/heltec/ui/core/backlight_policy.hpp", "src/heltec/ui/core/backlight_policy.cpp", "src/heltec/ui/screens/tracker_screen.hpp", "src/heltec/ui/screens/tracker_screen.cpp"]),
    (11, "指南针与查找朋友", 160, ["src/heltec/ui/screens/compass_ui_common.hpp", "src/heltec/ui/screens/compass_screen.hpp", "src/heltec/ui/screens/compass_screen.cpp", "src/heltec/ui/screens/find_friend_screen.hpp", "src/heltec/ui/screens/find_friend_screen_controls.cpp"]),
    (12, "地图与跟踪页面", 140, ["src/heltec/ui/map/map_panel_ids.hpp", "src/heltec/ui/map/map_panel.hpp", "src/heltec/ui/map/map_panel.cpp"]),
    (13, "系统设置与转发模式", 200, ["src/heltec/ui/screens/system_screen.hpp", "src/heltec/ui/screens/system_screen_rows.cpp", "src/heltec/ui/overlays/repeat_mode_overlay.hpp", "src/heltec/ui/overlays/repeat_mode_overlay.cpp"]),
    (14, "输入、导航与显示适配", 220, ["src/heltec/ui/core/input_types.hpp", "src/heltec/ui/core/input_adapter.cpp", "src/heltec/ui/core/input_dispatcher.cpp", "src/heltec/ui/core/input_pipeline.cpp", "src/heltec/ui/core/focus_key_mapper.cpp", "src/heltec/drivers/input/key_input.cpp", "src/heltec/drivers/input/momentary_button.cpp"]),
    (15, "连接方式与外部应用协同", 120, ["src/heltec/app/mesh_app_ui_link.cpp", "src/heltec/ui/widgets/top_pane.hpp", "src/heltec/ui/widgets/top_pane.cpp"]),
    (16, "数据保存与恢复", 140, ["src/heltec/ui/screens/system_screen_actions.cpp", "src/heltec/ui/core/ui_deferred_queue.hpp", "src/heltec/ui/core/ui_deferred_queue.cpp"]),
    (17, "当前版本功能变更", 120, ["src/heltec/ui/app/ui_app_navigation.cpp", "src/heltec/ui/theme/ui_theme_metrics.hpp", "src/heltec/ui/theme/ui_theme_metrics.cpp"]),
    (18, "编译、烧录与验收", 150, ["src/heltec/drivers/display/ssd1306_port.cpp", "src/heltec/drivers/display/st7735_port.cpp", "src/heltec/drivers/display/st7789_port.cpp"]),
    (19, "常见问题", 120, ["src/heltec/ui/overlays/alert_overlay.hpp", "src/heltec/ui/overlays/alert_overlay.cpp", "src/heltec/ui/core/ui_feedback.hpp", "src/heltec/ui/core/operation_hints.hpp"]),
    (20, "软件功能总结", 90, ["src/heltec/ui/theme/ui_theme_mono_flat.cpp", "src/heltec/ui/theme/ui_theme_pixel.cpp", "src/heltec/ui/theme/ui_widget_theme.hpp"]),
]

FORBIDDEN_PATH_PARTS = (
    "src/mesh.cpp", "src/packet", "src/identity", "datastore", "crypto",
    "ed25519", "license", "payload", "protocol", "transportkeystore",
    "geodesic", "map_geo", "heltecmesh.cpp",
)


def set_font(run, name="宋体", size=5, bold=False):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    for kind, value in (("begin", None), (None, " PAGE "), ("separate", None), ("end", None)):
        if value is not None:
            node = OxmlElement("w:instrText")
            node.set(qn("xml:space"), "preserve")
            node.text = value
        else:
            node = OxmlElement("w:fldChar")
            node.set(qn("w:fldCharType"), kind)
        run._r.append(node)
    set_font(run)


def description_for(path):
    low = path.lower()
    if "/screens/" in low:
        return "以下代码负责对应功能页面的控件创建、状态显示、焦点和用户操作。"
    if "/overlays/" in low:
        return "以下代码负责弹出提示、消息选择或临时操作界面的显示与交互。"
    if "/drivers/display/" in low:
        return "以下代码负责显示器初始化、LVGL绘制缓冲区和屏幕刷新适配。"
    if "/drivers/input/" in low or "/ui/core/input" in low:
        return "以下代码负责按键或触摸输入的读取、转换和界面事件分发。"
    if "/theme/" in low:
        return "以下代码负责字体、颜色、间距和控件视觉主题。"
    if "/map/" in low:
        return "以下代码负责地图面板的界面对象、可视区域和用户操作。"
    if "/widgets/" in low:
        return "以下代码负责通用界面控件的布局和状态刷新。"
    if "/app/" in low:
        return "以下代码负责界面应用层的页面组织和用户可见状态协调。"
    return "以下代码属于板级、界面或通用设备适配。"


def source_lines(path):
    lines = (ROOT / path).read_text(encoding="utf-8", errors="replace").splitlines()
    for line_no, line in enumerate(lines, 1):
        if not line.strip():
            continue
        if re.search(r"\b(author|created by|copyright)\b", line, re.I):
            continue
        if re.search(r"\b20\d{2}[-/.年]\d{1,2}", line):
            continue
        yield line_no, line.rstrip()


def collect_chapter(chapter_no, title, quota, files, globally_used):
    rows = [
        ("", 0, f"// ============================================================"),
        ("", 0, f"// 对应说明书第{chapter_no}章：{title}"),
        ("", 0, "// 本章源程序与软件说明书对应章节的功能及操作内容相对应。"),
    ]
    for rel in files:
        low = rel.replace("\\", "/").lower()
        if any(term in low for term in FORBIDDEN_PATH_PARTS):
            raise ValueError(f"章节{chapter_no}包含禁止文件：{rel}")
        path = ROOT / rel
        if not path.is_file():
            raise FileNotFoundError(path)
        rows.append((rel, 0, f"// 文件：{rel}"))
        rows.append((rel, 0, f"// 【功能注释】{description_for(rel)}"))
        emitted = 0
        for line_no, line in source_lines(rel):
            key = (rel, line_no)
            if key in globally_used:
                continue
            globally_used.add(key)
            if emitted > 0 and emitted % 70 == 0:
                rows.append((rel, 0, "// 【重要代码注释】以下部分继续处理当前模块的布局、状态或设备适配。"))
                if len(rows) >= quota:
                    return rows[:quota]
            rows.append((rel, line_no, line))
            emitted += 1
            if len(rows) >= quota:
                return rows[:quota]
    raise RuntimeError(f"第{chapter_no}章非核心代码不足：需要{quota}行，实际{len(rows)}行")


def add_code_line(doc, sequence, original_line, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = Pt(12)
    p.paragraph_format.widow_control = False
    # 鉴别材料中直接呈现源码，不添加文档序号或原始文件行号。
    # 英文代码使用 Times New Roman，中文注释及中文标点使用宋体。
    normalized = text.replace("\t", "  ")
    for part in CJK_PATTERN.split(normalized):
        if not part:
            continue
        is_chinese = bool(CJK_PATTERN.fullmatch(part))
        set_font(p.add_run(part), "宋体" if is_chinese else "Times New Roman", CODE_FONT_PT)


def build():
    if sum(item[2] for item in CHAPTERS) != TOTAL_LINES:
        raise RuntimeError("章节源码行数总和必须等于3000")

    used = set()
    selected = []
    chapter_offsets = []
    for number, title, quota, files in CHAPTERS:
        chapter_offsets.append((number, title, len(selected) + 1, len(selected) + quota))
        selected.extend(collect_chapter(number, title, quota, files, used))

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc = Document()
    cover = doc.sections[0]
    cover.top_margin = Cm(2.54)
    cover.bottom_margin = Cm(2.54)
    cover.left_margin = Cm(3.0)
    cover.right_margin = Cm(2.5)
    cover.header.paragraphs[0].text = ""
    cover.footer.paragraphs[0].text = ""

    for _ in range(3):
        doc.add_paragraph()
    for text, size, bold, after in (
        (NAME, 20, True, 28),
        ("源程序鉴别材料", 19, True, 50),
        ("编排方式：与软件说明书章节及顺序对应", 12, False, 10),
        ("代码范围：与软件说明书章节对应的源程序", 12, False, 10),
        ("源码正文：3000行，共60页，每页50行", 12, False, 10),
        ("字体：五号", 12, False, 10),
    ):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(after)
        set_font(p.add_run(text), "宋体", size, bold)

    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.top_margin = Cm(1.3)
    section.bottom_margin = Cm(1.3)
    section.left_margin = Cm(1.6)
    section.right_margin = Cm(1.4)
    section.header_distance = Cm(0.5)
    section.footer_distance = Cm(0.5)
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(header.add_run(NAME))
    section.footer.paragraphs[0].clear()
    add_page_number(section.footer.paragraphs[0])

    for index, (_, original_line, code) in enumerate(selected, 1):
        add_code_line(doc, index, original_line, code)
        if index % LINES_PER_PAGE == 0 and index != len(selected):
            doc.add_page_break()

    doc.core_properties.title = f"{NAME} 源程序鉴别材料"
    doc.core_properties.subject = "按软件说明书章节顺序编排的源程序"
    doc.core_properties.author = ""
    doc.core_properties.last_modified_by = ""
    doc.core_properties.comments = ""
    doc.save(OUT)

    print(f"DOCX={OUT}")
    print(f"CODE_LINES={len(selected)}")
    print(f"CODE_PAGES={len(selected) // LINES_PER_PAGE}")
    for number, title, start, end in chapter_offsets:
        print(f"CHAPTER_{number}={start}-{end} {title}")


if __name__ == "__main__":
    build()
